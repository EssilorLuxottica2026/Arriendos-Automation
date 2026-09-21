from pathlib import Path
from uuid import uuid4
import json
import re
import shutil
from datetime import datetime
from decimal import Decimal, InvalidOperation, ROUND_DOWN

from flask import (
    Flask,
    flash,
    redirect,
    render_template,
    request,
    send_file,
    url_for,
)
from werkzeug.utils import secure_filename
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .pipeline.pdf_reader import InvoiceSupportReader, NonInvoiceUBLDocument
from .pipeline.processor import LeaseAccountingPipeline, PipelineError, normalize_learning_phrase
from .runtime import (
    DATA_DIR,
    DICTIONARY_BACKUP_DIR,
    INPUT_DIR,
    LEARNING_DICTIONARY_PATH,
    OUTPUT_DIR,
    RAW_DIR,
    RESOURCE_DIR,
    UPLOAD_DIR,
)


BASE_DIR = RESOURCE_DIR
ALLOWED_EXTENSIONS = {".xlsx", ".xls", ".xlsm", ".csv", ".xml"}
SUPPORT_EXTENSIONS = {".pdf", ".xml"}


app = Flask(
    __name__,
    template_folder=str(BASE_DIR / "templates"),
    static_folder=str(BASE_DIR / "static"),
)
app.config["SECRET_KEY"] = "lease-accounting-local-dev"
app.config["MAX_CONTENT_LENGTH"] = 50 * 1024 * 1024

DATA_DIR.mkdir(exist_ok=True)
UPLOAD_DIR.mkdir(exist_ok=True)
OUTPUT_DIR.mkdir(exist_ok=True)
INPUT_DIR.mkdir(exist_ok=True)
RAW_DIR.mkdir(exist_ok=True)


# Standard filenames for reference databases
MASTER_FILENAMES = {
    "control_file": "CONTROL_ARRI_ADMON.xlsx",
    "contracts_file": "Contratos_con_condiciones.xlsx",
    "prorateo_file": "PRORATEO.xlsx",
    "distribution_file": "Cuadro_de_distribucion.xls",
    "history_file": "FACTURAS_CONTABILIZADAS.xlsx",
    "macro_workbook_file": "Arriendos_Macro.xlsm"
}
REQUIRED_MASTER_KEYS = (
    "control_file",
    "contracts_file",
    "prorateo_file",
    "distribution_file",
)


def _master_status() -> list[dict[str, str | bool]]:
    """Describe the locally saved databases; none are needed just to open the app."""
    return [
        {
            "key": key,
            "filename": filename,
            "available": (INPUT_DIR / filename).exists(),
            "required": key in REQUIRED_MASTER_KEYS,
        }
        for key, filename in MASTER_FILENAMES.items()
    ]


def _missing_required_masters(files) -> list[str]:
    """An upload satisfies the requirement even before it is saved locally."""
    return [
        MASTER_FILENAMES[key]
        for key in REQUIRED_MASTER_KEYS
        if not (
            (files.get(key) and files.get(key).filename)
            or (INPUT_DIR / MASTER_FILENAMES[key]).exists()
        )
    ]


def _allowed(filename: str, allowed_extensions: set[str] | None = None) -> bool:
    target_extensions = allowed_extensions or ALLOWED_EXTENSIONS
    return Path(filename).suffix.lower() in target_extensions


def _save_upload(file_storage, batch_dir: Path, allowed_extensions: set[str] | None = None) -> Path:
    filename = secure_filename(file_storage.filename or "")
    if not filename:
        raise PipelineError("One of the uploaded files does not have a valid filename.")
    if not _allowed(filename, allowed_extensions=allowed_extensions):
        allowed_hint = ", ".join(sorted(ext.lstrip(".") for ext in (allowed_extensions or ALLOWED_EXTENSIONS)))
        raise PipelineError(f"Unsupported file type for {filename}. Use {allowed_hint}.")

    target = batch_dir / filename
    file_storage.save(target)
    return target


def _resolve_and_save_master(key: str, file_storage, batch_dir: Path) -> Path | None:
    standard_name = MASTER_FILENAMES.get(key)
    if not standard_name:
        return None

    # If the user uploaded a new version of the file:
    if file_storage and file_storage.filename:
        temp_path = _save_upload(file_storage, batch_dir)
        target_master_path = INPUT_DIR / standard_name
        
        # If the file already exists in data/input, back it up to data/raw/
        if target_master_path.exists():
            timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
            stem = target_master_path.stem
            suffix = target_master_path.suffix
            backup_name = f"{stem}_replaced_{timestamp}{suffix}"
            
            backup_path = RAW_DIR / backup_name
            shutil.move(str(target_master_path), str(backup_path))
        
        # Copy the new file to data/input
        shutil.copy2(str(temp_path), str(target_master_path))
        return target_master_path

    # If the user did NOT upload a file, check if it exists in data/input/
    target_master_path = INPUT_DIR / standard_name
    if target_master_path.exists():
        return target_master_path

    # If it is a required master file but doesn't exist anywhere
    if key in ["control_file", "contracts_file", "prorateo_file", "distribution_file"]:
        raise PipelineError(f"El archivo maestro obligatorio '{standard_name}' no se encuentra en el servidor. Por favor, súbelo al menos una vez.")
        
    return None

def _resolve_existing_master(key: str) -> Path | None:
    standard_name = MASTER_FILENAMES.get(key)
    if not standard_name:
        return None

    target_master_path = INPUT_DIR / standard_name
    if target_master_path.exists():
        return target_master_path
    if key in REQUIRED_MASTER_KEYS:
        raise PipelineError(
            f"El archivo maestro obligatorio '{standard_name}' no se encuentra en el servidor. "
            "Guárdalo desde la sección 2 antes de procesar."
        )
    return None


def _job_path(batch_id: str) -> Path:
    return UPLOAD_DIR / batch_id / "concept_review_job.json"


def _save_review_job(batch_id: str, job: dict) -> None:
    path = _job_path(batch_id)
    path.write_text(json.dumps(job, ensure_ascii=False, indent=2), encoding="utf-8")


def _load_review_job(batch_id: str) -> dict:
    path = _job_path(batch_id)
    if not path.exists():
        raise PipelineError("No se encontro la sesion de revision de conceptos. Vuelve a procesar el archivo.")
    return json.loads(path.read_text(encoding="utf-8"))


def _equal_split_percentages(split_factor: int) -> list[str]:
    if split_factor <= 1:
        return ["100"]
    base = (Decimal("100") / Decimal(split_factor)).quantize(
        Decimal("0.01"),
        rounding=ROUND_DOWN,
    )
    percentages = [base for _ in range(split_factor)]
    percentages[-1] += Decimal("100") - sum(percentages)
    return [format(value, "f") for value in percentages]


def _build_split_review_items(job: dict) -> list[dict]:
    reader = InvoiceSupportReader()
    items = []
    for support_path, raw_factor in (job.get("support_split_factors") or {}).items():
        split_factor = max(int(raw_factor or 1), 1)
        if split_factor <= 1:
            continue
        path = Path(support_path)
        parsed = reader.parse_file(path)
        total = parsed.total if parsed.total is not None else parsed.subtotal
        items.append(
            {
                "support_path": support_path,
                "source_file": path.name,
                "invoice_id": parsed.invoice_id or "Sin numero identificado",
                "supplier_name": parsed.supplier_name or "Proveedor no identificado",
                "split_factor": split_factor,
                "total": total,
                "equal_percentages": _equal_split_percentages(split_factor),
            }
        )
    return items


def _render_split_review(batch_id: str, job: dict):
    items = _build_split_review_items(job)
    _save_review_job(batch_id, job)
    return render_template(
        "split_review.html",
        batch_id=batch_id,
        items=items,
        manual_processing_files=job.get("manual_processing_files", []),
    )


def _run_pipeline_from_job(job: dict) -> dict:
    pipeline = LeaseAccountingPipeline(output_dir=OUTPUT_DIR, learning_dictionary_path=LEARNING_DICTIONARY_PATH)
    result = pipeline.run(
        invoices_path=Path(job["invoices_path"]),
        control_path=Path(job["control_path"]),
        contracts_path=Path(job["contracts_path"]),
        prorateo_path=Path(job["prorateo_path"]),
        distribution_path=Path(job["distribution_path"]),
        history_path=Path(job["history_path"]) if job.get("history_path") else None,
        support_paths=[Path(path) for path in job.get("support_paths", [])],
        support_split_factors=job.get("support_split_factors") or None,
        support_split_weights=job.get("support_split_weights") or None,
        macro_template_path=Path(job["macro_template_path"]) if job.get("macro_template_path") else None,
        period=job.get("period") or None,
        contract_selections=job.get("contract_selections") or None,
        discount_selections=job.get("discount_selections") or None,
        beneficiary_selections=job.get("beneficiary_selections") or None,
    )
    manual_processing_files = job.get("manual_processing_files", [])
    result["manual_processing_files"] = manual_processing_files
    result.setdefault("processing_issues", [])
    result["processing_issues"].extend(
        {
            "invoice_id": "No identificada",
            "source_file": item.get("filename"),
            "store": None,
            "location": item.get("filename") or "Archivo UBL sin nombre",
            "problem": (
                f"Es un documento UBL {item.get('document_type') or 'no identificado'}, "
                "no una factura XML. Debe procesarse manualmente."
            ),
            "possible_solution": (
                "Solicita al proveedor el XML de la factura o nota de credito correspondiente, "
                "o procesa este documento manualmente."
            ),
            "issue_type": "non_invoice_ubl",
        }
        for item in manual_processing_files
    )
    _write_processing_issues_report(result)
    return result


def _safe_excel_text(value) -> str:
    text = str(value or "").strip()
    if text.startswith(("=", "+", "-", "@")):
        return f"'{text}"
    return text


def _write_processing_issues_report(result: dict) -> None:
    issues = result.get("processing_issues") or []
    result["omitted_invoices_report"] = None
    if not issues:
        return

    try:
        output_reference = Path(result.get("output_csv") or "")
        run_folder = Path(result.get("run_folder") or output_reference.parent)
        if str(run_folder) in {"", "."}:
            raise ValueError("No se pudo identificar la carpeta de resultados del lote.")

        report_name = "facturas_omitidas.xlsx"
        report_dir = OUTPUT_DIR / run_folder
        report_dir.mkdir(parents=True, exist_ok=True)
        report_path = report_dir / report_name

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Facturas omitidas"
        headers = ["Factura", "Archivo", "Tienda", "Razon", "Posible solucion"]
        sheet.append(headers)
        for issue in issues:
            sheet.append(
                [
                    _safe_excel_text(issue.get("invoice_id") or "Sin numero"),
                    _safe_excel_text(issue.get("source_file") or ""),
                    _safe_excel_text(issue.get("store") or ""),
                    _safe_excel_text(issue.get("problem") or "Requiere revision manual"),
                    _safe_excel_text(
                        issue.get("possible_solution")
                        or "Revisa la factura y sus datos en los archivos maestros antes de reprocesarla."
                    ),
                ]
            )

        header_fill = PatternFill("solid", fgColor="152134")
        header_font = Font(color="FFFFFF", bold=True)
        for cell in sheet[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(vertical="center")

        widths = [18, 42, 16, 62, 70]
        for column_index, width in enumerate(widths, start=1):
            sheet.column_dimensions[get_column_letter(column_index)].width = width
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)

        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = f"A1:E{sheet.max_row}"
        sheet.row_dimensions[1].height = 24
        workbook.save(report_path)
        result["omitted_invoices_report"] = f"{run_folder.as_posix()}/{report_name}"
    except Exception as exc:
        result.setdefault("warnings", []).append(
            f"No se pudo generar el Excel de facturas omitidas: {exc}"
        )


def _render_all_beneficiary_invoices_skipped(job: dict, review_items: list[dict]):
    issues = []
    selections = job.get("beneficiary_selections") or {}
    for item in review_items:
        if selections.get(item.get("review_key")) != "skip":
            continue
        total = item.get("items_total") or 0
        issues.append(
            {
                "invoice_id": item.get("invoice_id") or "Sin numero",
                "source_file": item.get("source_file"),
                "store": item.get("store"),
                "location": f"{item.get('source_file') or ''} / {item.get('store') or ''}",
                "problem": (
                    f"El XML contiene {item.get('line_count') or 0} lineas iguales asignadas a "
                    f"beneficiarios diferentes por un total de {total}; la factura fue omitida "
                    "por decision del usuario"
                ),
                "possible_solution": (
                    f"Revisa en el PDF si el concepto aparece como una sola linea por {total}. "
                    "Luego vuelve a procesar y elige consolidar o continuar normalmente."
                ),
                "issue_type": "beneficiary_distribution_skipped",
            }
        )

    for item in job.get("manual_processing_files", []):
        issues.append(
            {
                "invoice_id": "No identificada",
                "source_file": item.get("filename"),
                "store": None,
                "location": item.get("filename") or "Archivo UBL sin nombre",
                "problem": (
                    f"Es un documento UBL {item.get('document_type') or 'no identificado'}, "
                    "no una factura XML. Debe procesarse manualmente."
                ),
                "possible_solution": (
                    "Solicita al proveedor el XML de la factura o nota de credito correspondiente, "
                    "o procesa este documento manualmente."
                ),
                "issue_type": "non_invoice_ubl",
            }
        )

    run_folder = datetime.now().strftime("%Y.%m.%d_%H.%M.%S")
    result = {
        "run_folder": run_folder,
        "output_csv": None,
        "invoice_csv_bundle": None,
        "manual_style_csv_bundle": None,
        "output_rows": 0,
        "header_rows": 0,
        "validation_rows": len(issues),
        "warnings": [
            "Todas las facturas procesables del lote fueron omitidas por decision del usuario."
        ],
        "logs": ["Proceso completado sin generar CSV: todas las facturas fueron omitidas."],
        "processing_issues": issues,
    }
    _write_processing_issues_report(result)
    return render_template("result.html", result=result)


def _prepare_pipeline_from_job(job: dict):
    pipeline = LeaseAccountingPipeline(output_dir=OUTPUT_DIR, learning_dictionary_path=LEARNING_DICTIONARY_PATH)
    data, output_df, validation_df = pipeline.prepare(
        invoices_path=Path(job["invoices_path"]),
        control_path=Path(job["control_path"]),
        contracts_path=Path(job["contracts_path"]),
        prorateo_path=Path(job["prorateo_path"]),
        distribution_path=Path(job["distribution_path"]),
        history_path=Path(job["history_path"]) if job.get("history_path") else None,
        support_paths=[Path(path) for path in job.get("support_paths", [])],
        support_split_factors=job.get("support_split_factors") or None,
        support_split_weights=job.get("support_split_weights") or None,
        macro_template_path=Path(job["macro_template_path"]) if job.get("macro_template_path") else None,
        period=job.get("period") or None,
        contract_selections=job.get("contract_selections") or None,
        discount_selections=job.get("discount_selections") or None,
        beneficiary_selections=job.get("beneficiary_selections") or None,
    )
    return pipeline, data, output_df, validation_df


def _continue_processing_job(batch_id: str, job: dict):
    pipeline, _, output_df, _ = _prepare_pipeline_from_job(job)
    manual_processing_files = job.get("manual_processing_files", [])
    if pipeline.beneficiary_review_items:
        job["beneficiary_review_context"] = pipeline.beneficiary_review_items
        _save_review_job(batch_id, job)
        return render_template(
            "beneficiary_review.html",
            batch_id=batch_id,
            items=pipeline.beneficiary_review_items,
            warnings=pipeline.warnings,
            manual_processing_files=manual_processing_files,
        )

    if pipeline.discount_review_items:
        _save_review_job(batch_id, job)
        return render_template(
            "discount_review.html",
            batch_id=batch_id,
            items=pipeline.discount_review_items,
            warnings=pipeline.warnings,
            manual_processing_files=manual_processing_files,
        )

    if pipeline.contract_review_items:
        _save_review_job(batch_id, job)
        return render_template(
            "contract_review.html",
            batch_id=batch_id,
            items=pipeline.contract_review_items,
            warnings=pipeline.warnings,
            manual_processing_files=manual_processing_files,
        )

    review_items = pipeline.build_concept_review_items(output_df)
    if review_items:
        _save_review_job(batch_id, job)
        return render_template(
            "concept_review.html",
            batch_id=batch_id,
            items=review_items,
            concepts=pipeline.learning_concepts(),
            warnings=pipeline.warnings,
            manual_processing_files=manual_processing_files,
        )

    result = _run_pipeline_from_job(job)
    return render_template("result.html", result=result)


def _separate_non_invoice_xmls(
    support_paths: list[Path],
    display_names: dict[str, str] | None = None,
) -> tuple[list[Path], list[dict]]:
    valid_paths = []
    manual_files = []
    reader = InvoiceSupportReader()
    for path in support_paths:
        if path.suffix.lower() != ".xml":
            valid_paths.append(path)
            continue
        try:
            reader.parse_file(path)
            valid_paths.append(path)
        except NonInvoiceUBLDocument as exc:
            manual_files.append(
                {
                    "filename": (display_names or {}).get(str(path), exc.filename),
                    "document_type": exc.document_type,
                }
            )
    return valid_paths, manual_files


def _append_learning_dictionary_phrases(assignments: list[dict]) -> None:
    if not LEARNING_DICTIONARY_PATH.exists():
        raise PipelineError(f"No existe el diccionario de conceptos: {LEARNING_DICTIONARY_PATH.name}")
    try:
        workbook = load_workbook(LEARNING_DICTIONARY_PATH)
    except PermissionError as exc:
        raise PipelineError("Cierra el archivo Diccionario_Conceptos_Simple.xlsx antes de guardar nuevas reglas.") from exc
    if "Diccionario" not in workbook.sheetnames:
        raise PipelineError("El diccionario debe tener una hoja llamada 'Diccionario'.")

    sheet = workbook["Diccionario"]
    headers = {str(cell.value or "").strip().lower(): idx for idx, cell in enumerate(sheet[1], start=1)}
    concept_col = headers.get("concepto")
    phrases_col = headers.get("frases/palabras") or headers.get("frases") or headers.get("palabras")
    if not concept_col or not phrases_col:
        raise PipelineError("El diccionario debe tener columnas 'Concepto' y 'Frases/palabras'.")

    concept_rows = {}
    for row_idx in range(2, sheet.max_row + 1):
        concept = str(sheet.cell(row_idx, concept_col).value or "").strip()
        if concept:
            concept_rows[concept.upper()] = row_idx

    for assignment in assignments:
        concept = (assignment.get("concept") or "").strip()
        phrase_text = (assignment.get("phrase") or "").strip()
        if not concept or not phrase_text:
            continue
        row_idx = concept_rows.get(concept.upper())
        if not row_idx:
            row_idx = sheet.max_row + 1
            sheet.cell(row_idx, concept_col).value = concept
            concept_rows[concept.upper()] = row_idx
        current = str(sheet.cell(row_idx, phrases_col).value or "").strip()
        existing = []
        existing_keys = set()
        for part in current.split(","):
            normalized = normalize_learning_phrase(part)
            normalized_key = normalized.upper()
            if normalized and normalized_key not in existing_keys:
                existing.append(normalized)
                existing_keys.add(normalized_key)
        for phrase in [part.strip() for part in phrase_text.split(",") if part.strip()]:
            phrase = normalize_learning_phrase(phrase)
            if phrase and phrase.upper() not in existing_keys:
                existing.append(phrase)
                existing_keys.add(phrase.upper())
        sheet.cell(row_idx, phrases_col).value = ", ".join(existing)

    _backup_learning_dictionary()
    try:
        workbook.save(LEARNING_DICTIONARY_PATH)
    except PermissionError as exc:
        raise PipelineError(
            "No pude guardar el diccionario. Cierra Diccionario_Conceptos_Simple.xlsx "
            "y comprueba que la carpeta de la aplicacion tenga permisos de escritura."
        ) from exc


@app.route("/", methods=["GET"])
def index():
    return render_template("index.html", master_status=_master_status())


def _dictionary_sheet(workbook):
    if "Diccionario" not in workbook.sheetnames:
        raise PipelineError("El diccionario debe tener una hoja llamada 'Diccionario'.")
    sheet = workbook["Diccionario"]
    headers = {str(cell.value or "").strip().lower(): idx for idx, cell in enumerate(sheet[1], start=1)}
    concept_col = headers.get("concepto")
    account_col = headers.get("gl account") or headers.get("gl_account") or headers.get("cuenta")
    phrases_col = headers.get("frases/palabras") or headers.get("frases") or headers.get("palabras")
    sigla_col = headers.get("sigla") or headers.get("texto_corto_lucy") or headers.get("texto")
    if not concept_col or not account_col or not phrases_col or not sigla_col:
        raise PipelineError("El diccionario debe tener columnas 'Concepto', 'GL Account', 'Frases/palabras' y 'Sigla'.")
    return sheet, concept_col, account_col, phrases_col, sigla_col


def _load_dictionary_workbook():
    if not LEARNING_DICTIONARY_PATH.exists():
        raise PipelineError(f"No existe el diccionario de conceptos: {LEARNING_DICTIONARY_PATH.name}")
    try:
        return load_workbook(LEARNING_DICTIONARY_PATH)
    except PermissionError as exc:
        raise PipelineError("Cierra Diccionario_Conceptos_Simple.xlsx para modificar la configuracion.") from exc


def _backup_learning_dictionary() -> Path:
    if not LEARNING_DICTIONARY_PATH.exists():
        raise PipelineError(f"No existe el diccionario de conceptos: {LEARNING_DICTIONARY_PATH.name}")
    DICTIONARY_BACKUP_DIR.mkdir(parents=True, exist_ok=True)
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S_%f")
    backup_path = DICTIONARY_BACKUP_DIR / f"{LEARNING_DICTIONARY_PATH.stem}_{timestamp}.xlsx"
    try:
        shutil.copy2(LEARNING_DICTIONARY_PATH, backup_path)
    except PermissionError as exc:
        raise PipelineError(
            "No pude respaldar el diccionario. Cierra Diccionario_Conceptos_Simple.xlsx "
            "y comprueba los permisos de la carpeta."
        ) from exc

    backups = sorted(
        DICTIONARY_BACKUP_DIR.glob(f"{LEARNING_DICTIONARY_PATH.stem}_*.xlsx"),
        key=lambda path: path.stat().st_mtime,
        reverse=True,
    )
    for old_backup in backups[20:]:
        old_backup.unlink(missing_ok=True)
    return backup_path


def _validate_concept_account(concept: str, account: str, sigla: str) -> tuple[str, str, str]:
    concept = re.sub(r"\s+", " ", concept).strip().upper()
    account = re.sub(r"\s+", "_", account).strip().upper()
    sigla = re.sub(r"\s+", " ", sigla).strip().upper() if sigla else ""
    if not concept or not account:
        raise PipelineError("El concepto y el GL Account son obligatorios.")
    if account != "SEGUN_CONTRATO" and not re.fullmatch(r"\d{6,15}", account):
        raise PipelineError("El GL Account debe contener entre 6 y 15 digitos, o usar SEGUN_CONTRATO para RENTA.")
    return concept, account, sigla


@app.route("/settings", methods=["GET", "POST"])
def settings():
    try:
        workbook = _load_dictionary_workbook()
        sheet, concept_col, account_col, phrases_col, sigla_col = _dictionary_sheet(workbook)

        if request.method == "POST":
            action = request.form.get("action", "").strip().lower()
            rows_by_concept = {
                str(sheet.cell(row_idx, concept_col).value or "").strip().upper(): row_idx
                for row_idx in range(2, sheet.max_row + 1)
                if str(sheet.cell(row_idx, concept_col).value or "").strip()
            }

            if action == "delete":
                original = request.form.get("original_concept", "").strip().upper()
                row_idx = rows_by_concept.get(original)
                if not row_idx:
                    raise PipelineError("El concepto que intentas eliminar ya no existe.")
                sheet.delete_rows(row_idx, 1)
                success_message = f"Concepto {original} eliminado correctamente."
            else:
                concept, account, sigla = _validate_concept_account(
                    request.form.get("concept", ""),
                    request.form.get("account", ""),
                    request.form.get("sigla", ""),
                )

            if action == "add":
                if concept in rows_by_concept:
                    raise PipelineError(f"El concepto {concept} ya existe. Puedes editarlo en la tabla.")
                row_idx = sheet.max_row + 1
                sheet.cell(row_idx, concept_col).value = concept
                sheet.cell(row_idx, account_col).value = account
                sheet.cell(row_idx, sigla_col).value = sigla
                sheet.cell(row_idx, phrases_col).value = ""
                success_message = f"Concepto {concept} agregado correctamente."
            elif action == "update":
                original = request.form.get("original_concept", "").strip().upper()
                row_idx = rows_by_concept.get(original)
                if not row_idx:
                    raise PipelineError("El concepto que intentas editar ya no existe.")
                duplicate_row = rows_by_concept.get(concept)
                if duplicate_row and duplicate_row != row_idx:
                    raise PipelineError(f"Ya existe otro concepto llamado {concept}.")
                sheet.cell(row_idx, concept_col).value = concept
                sheet.cell(row_idx, account_col).value = account
                sheet.cell(row_idx, sigla_col).value = sigla
                success_message = f"Concepto {concept} actualizado correctamente."
            elif action != "delete":
                raise PipelineError("Accion de configuracion no valida.")

            _backup_learning_dictionary()
            try:
                workbook.save(LEARNING_DICTIONARY_PATH)
            except PermissionError as exc:
                raise PipelineError(
                    "Cierra Diccionario_Conceptos_Simple.xlsx y comprueba que la carpeta "
                    "de la aplicacion tenga permisos de escritura."
                ) from exc
            flash(success_message, "success")
            return redirect(url_for("settings"))

        concepts = []
        for row_idx in range(2, sheet.max_row + 1):
            concept = str(sheet.cell(row_idx, concept_col).value or "").strip()
            if not concept:
                continue
            phrases = [part.strip() for part in str(sheet.cell(row_idx, phrases_col).value or "").split(",") if part.strip()]
            concepts.append(
                {
                    "concept": concept,
                    "account": str(sheet.cell(row_idx, account_col).value or "").strip(),
                    "sigla": str(sheet.cell(row_idx, sigla_col).value or "").strip(),
                        "phrase_count": len(phrases),
                }
            )
        return render_template("settings.html", concepts=concepts)
    except PipelineError as exc:
        flash(str(exc), "error")
        return redirect(url_for("settings")) if request.method == "POST" else redirect(url_for("index"))


@app.route("/process", methods=["POST"])
def process():
    if not LEARNING_DICTIONARY_PATH.exists():
        flash(
            "No se encontro Diccionario_Conceptos_Simple.xlsx junto a la aplicacion. "
            "Colocalo en la misma carpeta que el ejecutable antes de procesar."
        )
        return redirect(url_for("index"))

    invoices_file = request.files.get("invoices_file")
    if not invoices_file or not invoices_file.filename:
        flash("Por favor, sube el archivo de facturas consolidado (Invoices file).")
        return redirect(url_for("index"))

    missing_masters = _missing_required_masters(request.files)
    if missing_masters:
        flash(
            "Para procesar el archivo del mes, sube por primera vez en la sección 2 las bases requeridas: "
            + ", ".join(missing_masters)
            + "."
        )
        return redirect(url_for("index"))

    batch_id = uuid4().hex
    batch_dir = UPLOAD_DIR / batch_id
    batch_dir.mkdir(parents=True, exist_ok=True)

    try:
        invoices_path = _save_upload(invoices_file, batch_dir)
        
        control_path = _resolve_existing_master("control_file")
        contracts_path = _resolve_existing_master("contracts_file")
        prorateo_path = _resolve_existing_master("prorateo_file")
        distribution_path = _resolve_existing_master("distribution_file")
        history_path = _resolve_existing_master("history_file")
        macro_workbook_path = _resolve_existing_master("macro_workbook_file")

        support_paths = []
        support_display_names = {}
        support_split_factors = {}
        for support_file in request.files.getlist("invoice_support_files"):
            if support_file and support_file.filename:
                support_path = _save_upload(support_file, batch_dir, allowed_extensions=SUPPORT_EXTENSIONS)
                support_paths.append(support_path)
                support_display_names[str(support_path)] = support_file.filename
        for field_name, split_factor in [("invoice_support_files_split_2", 2), ("invoice_support_files_split_3", 3)]:
            for support_file in request.files.getlist(field_name):
                if support_file and support_file.filename:
                    support_path = _save_upload(support_file, batch_dir, allowed_extensions=SUPPORT_EXTENSIONS)
                    support_paths.append(support_path)
                    support_display_names[str(support_path)] = support_file.filename
                    support_split_factors[str(support_path)] = split_factor

        support_paths, manual_processing_files = _separate_non_invoice_xmls(
            support_paths,
            display_names=support_display_names,
        )
        valid_support_keys = {str(path) for path in support_paths}
        support_split_factors = {
            path: factor
            for path, factor in support_split_factors.items()
            if path in valid_support_keys
        }

        if manual_processing_files and not support_paths:
            return render_template(
                "index.html",
                manual_processing_files=manual_processing_files,
            )

        period = request.form.get("period", "").strip()
        job = {
            "invoices_path": str(invoices_path),
            "control_path": str(control_path),
            "contracts_path": str(contracts_path),
            "prorateo_path": str(prorateo_path),
            "distribution_path": str(distribution_path),
            "history_path": str(history_path) if history_path else None,
            "support_paths": [str(path) for path in support_paths],
            "support_split_factors": support_split_factors,
            "support_split_weights": {},
            "macro_template_path": str(macro_workbook_path) if macro_workbook_path else None,
            "period": period or None,
            "manual_processing_files": manual_processing_files,
            "contract_selections": {},
            "discount_selections": {},
            "beneficiary_selections": {},
        }
        if support_split_factors:
            return _render_split_review(batch_id, job)
        return _continue_processing_job(batch_id, job)
    except PipelineError as exc:
        if "Todas las facturas del lote se marcaron para procesamiento manual" in str(exc):
            flash(str(exc), "warning")
        else:
            flash(str(exc))
        return redirect(url_for("index"))
    except Exception as exc:  # pragma: no cover
        flash(f"Unexpected error: {exc}")
        return redirect(url_for("index"))


@app.route("/save-masters", methods=["POST"])
def save_masters():
    master_keys = tuple(MASTER_FILENAMES)
    uploaded_keys = [
        key for key in master_keys
        if request.files.get(key) and request.files[key].filename
    ]
    if not uploaded_keys:
        flash("Selecciona al menos una base de datos en la sección 2 para guardarla.", "warning")
        return redirect(url_for("index"))

    batch_id = uuid4().hex
    batch_dir = UPLOAD_DIR / batch_id
    batch_dir.mkdir(parents=True, exist_ok=True)
    try:
        for key in uploaded_keys:
            _resolve_and_save_master(key, request.files.get(key), batch_dir)
        flash("Las bases de datos seleccionadas se guardaron correctamente.", "success")
    except PipelineError as exc:
        flash(str(exc), "error")
    except Exception as exc:  # pragma: no cover
        flash(f"Unexpected error: {exc}", "error")
    return redirect(url_for("index"))


@app.route("/split-review/<batch_id>", methods=["POST"])
def split_review(batch_id: str):
    try:
        job = _load_review_job(batch_id)
        split_factors = job.get("support_split_factors") or {}
        item_count = int(request.form.get("item_count", "0"))
        if item_count != len(split_factors):
            raise PipelineError("La revision de porcentajes no coincide con las facturas cargadas.")

        split_weights = {}
        for idx, (support_path, raw_factor) in enumerate(split_factors.items()):
            split_factor = max(int(raw_factor or 1), 1)
            percentages = []
            for part_index in range(split_factor):
                raw_value = request.form.get(
                    f"percentage_{idx}_{part_index}",
                    "",
                ).strip()
                try:
                    percentage = Decimal(raw_value)
                except (InvalidOperation, ValueError):
                    raise PipelineError(
                        f"Escribe porcentajes validos para {Path(support_path).name}."
                    )
                if not percentage.is_finite() or percentage <= 0:
                    raise PipelineError(
                        f"Cada porcentaje de {Path(support_path).name} debe ser mayor que cero."
                    )
                percentages.append(percentage)

            if sum(percentages) != Decimal("100"):
                raise PipelineError(
                    f"Los porcentajes de {Path(support_path).name} deben sumar exactamente 100%."
                )
            split_weights[support_path] = [format(value, "f") for value in percentages]

        job["support_split_weights"] = split_weights
        _save_review_job(batch_id, job)
        return _continue_processing_job(batch_id, job)
    except PipelineError as exc:
        flash(str(exc))
        try:
            job = _load_review_job(batch_id)
            return _render_split_review(batch_id, job)
        except PipelineError:
            return redirect(url_for("index"))
    except Exception as exc:  # pragma: no cover
        flash(f"Unexpected error: {exc}")
        return redirect(url_for("index"))


@app.route("/beneficiary-review/<batch_id>", methods=["POST"])
def beneficiary_review(batch_id: str):
    expected_items = []
    try:
        job = _load_review_job(batch_id)
        item_count = int(request.form.get("item_count", "0"))
        expected_items = job.get("beneficiary_review_context") or []
        if item_count != len(expected_items):
            raise PipelineError("La revision de beneficiarios no coincide con las facturas cargadas.")

        selections = dict(job.get("beneficiary_selections") or {})
        valid_actions = {"keep", "consolidate", "skip"}
        for idx in range(item_count):
            review_key = request.form.get(f"review_key_{idx}", "").strip()
            action = request.form.get(f"beneficiary_action_{idx}", "").strip().lower()
            if not review_key or action not in valid_actions:
                raise PipelineError("Selecciona una accion para cada distribucion por beneficiarios.")
            selections[review_key] = action

        job["beneficiary_selections"] = selections
        _save_review_job(batch_id, job)
        return _continue_processing_job(batch_id, job)
    except PipelineError as exc:
        if "Todas las facturas del lote se marcaron para procesamiento manual" in str(exc):
            return _render_all_beneficiary_invoices_skipped(job, expected_items)
        else:
            flash(str(exc))
        return redirect(url_for("index"))
    except Exception as exc:  # pragma: no cover
        flash(f"Unexpected error: {exc}")
        return redirect(url_for("index"))


@app.route("/contract-review/<batch_id>", methods=["POST"])
def contract_review(batch_id: str):
    try:
        job = _load_review_job(batch_id)
        item_count = int(request.form.get("item_count", "0"))
        selections = dict(job.get("contract_selections") or {})
        for idx in range(item_count):
            ceco = request.form.get(f"ceco_{idx}", "").strip()
            contract_key = request.form.get(f"contract_{idx}", "").strip()
            if not ceco or not contract_key:
                raise PipelineError("Selecciona un contrato para cada CECO con estados contradictorios.")
            selections[ceco] = contract_key
        job["contract_selections"] = selections
        _save_review_job(batch_id, job)
        return _continue_processing_job(batch_id, job)
    except PipelineError as exc:
        flash(str(exc))
        return redirect(url_for("index"))
    except Exception as exc:  # pragma: no cover
        flash(f"Unexpected error: {exc}")
        return redirect(url_for("index"))


@app.route("/discount-review/<batch_id>", methods=["POST"])
def discount_review(batch_id: str):
    try:
        job = _load_review_job(batch_id)
        item_count = int(request.form.get("item_count", "0"))
        selections = dict(job.get("discount_selections") or {})
        valid_actions = {"apply", "omit", "skip"}
        for idx in range(item_count):
            review_key = request.form.get(f"review_key_{idx}", "").strip()
            action = request.form.get(f"discount_action_{idx}", "").strip().lower()
            if not review_key or action not in valid_actions:
                raise PipelineError("Selecciona una accion para cada descuento que requiere revision.")
            selections[review_key] = action
        job["discount_selections"] = selections
        _save_review_job(batch_id, job)
        return _continue_processing_job(batch_id, job)
    except PipelineError as exc:
        if "Todas las facturas del lote se marcaron para procesamiento manual" in str(exc):
            flash(str(exc), "warning")
        else:
            flash(str(exc))
        return redirect(url_for("index"))
    except Exception as exc:  # pragma: no cover
        flash(f"Unexpected error: {exc}")
        return redirect(url_for("index"))


@app.route("/concept-review/<batch_id>", methods=["POST"])
def concept_review(batch_id: str):
    try:
        job = _load_review_job(batch_id)
        item_count = int(request.form.get("item_count", "0"))
        assignments = []
        for idx in range(item_count):
            concept = request.form.get(f"concept_{idx}", "").strip()
            phrase = request.form.get(f"phrase_{idx}", "").strip()
            if not concept or not phrase:
                raise PipelineError("Todos los items pendientes necesitan concepto y frase/palabra.")
            assignments.append({"concept": concept, "phrase": phrase})
        _append_learning_dictionary_phrases(assignments)
        result = _run_pipeline_from_job(job)
    except PipelineError as exc:
        flash(str(exc))
        return redirect(url_for("index"))
    except Exception as exc:  # pragma: no cover
        flash(f"Unexpected error: {exc}")
        return redirect(url_for("index"))

    return render_template("result.html", result=result)


@app.route("/download/<path:filename>", methods=["GET"])
def download(filename: str):
    file_path = OUTPUT_DIR / filename
    if not file_path.exists():
        flash("Requested file does not exist.")
        return redirect(url_for("index"))
    return send_file(file_path, as_attachment=True)
