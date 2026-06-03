from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


ROOT_DIR = Path(__file__).resolve().parents[1]
OUTPUT_PATH = ROOT_DIR / "docs" / "business" / "business_validation_template.xlsx"

HEADER_FILL = PatternFill(fill_type="solid", fgColor="1F2937")
HEADER_FONT = Font(color="FFFFFF", bold=True)
TITLE_FILL = PatternFill(fill_type="solid", fgColor="DCE6F2")
TITLE_FONT = Font(color="111827", bold=True)
WRAP = Alignment(wrap_text=True, vertical="top")


def write_table(ws, start_row, columns, rows, title=None):
    row_idx = start_row
    if title:
        ws.cell(row=row_idx, column=1, value=title)
        ws.cell(row=row_idx, column=1).fill = TITLE_FILL
        ws.cell(row=row_idx, column=1).font = TITLE_FONT
        row_idx += 1

    for col_idx, value in enumerate(columns, start=1):
        cell = ws.cell(row=row_idx, column=col_idx, value=value)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = WRAP

    row_idx += 1
    for row in rows:
        for col_idx, value in enumerate(row, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.alignment = WRAP
        row_idx += 1
    return row_idx


def autosize(ws, widths=None):
    widths = widths or {}
    for idx, column_cells in enumerate(ws.columns, start=1):
        letter = get_column_letter(idx)
        if letter in widths:
            ws.column_dimensions[letter].width = widths[letter]
            continue
        values = [str(cell.value) for cell in column_cells if cell.value not in (None, "")]
        if values:
            ws.column_dimensions[letter].width = min(max(len(v) for v in values) + 2, 40)


def build_instructions(ws):
    ws.title = "01_Instrucciones"
    rows = [
        ["Objetivo", "Confirmar las reglas reales del proceso para que la automatizacion genere archivos 100% confiables para Lucy y para el bot."],
        ["Quien debe llenarlo", "El dueno del proceso o la persona que conoce el SOP completo y que hoy realiza la preparacion manual en Lucy/SAP."],
        ["Como responder", "Usa lenguaje operativo. Si una regla depende del tipo de factura, anotalo. Si un campo no se usa, escribe NO APLICA."],
        ["Que es mejor", "Dejar un campo vacio a inventar una regla. Si no hay una fuente real, escribe que hoy se revisa manualmente."],
        ["Prioridad de llenado", "1) Casos reales 2) Header fields 3) Allocated costs 4) Splits 5) Reglas de bloqueo."],
        ["Archivos utiles para responder", "SOP, video, CSV real de Lucy, base macro, invoices file, PDFs reales, archivo de distribucion, prorrateo, contratos e historico."],
    ]
    write_table(ws, 1, ["Tema", "Respuesta esperada"], rows, title="Como usar esta plantilla")
    ws.freeze_panes = "A3"
    autosize(ws, {"A": 24, "B": 110})


def build_header_fields(ws):
    rows = [
        ["Invoice Date", "Si", "Si", "PDF / tabla / otro", "Fecha exacta que se pone en Lucy", "Ej. usar fecha de factura del PDF", "Si falta, que se hace?", ""],
        ["Invoice Number", "Si", "Si", "Tabla / PDF", "Numero exacto de factura", "", "", ""],
        ["Original Invoice Ref Number", "Si/No", "No", "Confirmar", "Si es igual al invoice number o lleva otra referencia", "", "", ""],
        ["Document Type Sap", "Si", "Si", "Regla fija o maestro", "Ej. KR", "", "", ""],
        ["Currency", "Si", "Si", "Regla fija o PDF", "Ej. COP", "", "", ""],
        ["Invoice Total Amount", "Si", "Si", "PDF / tabla", "Valor total exacto que se registra", "", "", ""],
        ["Total Taxable Amount", "Si", "Si", "PDF / calculo", "Base sin IVA", "", "", ""],
        ["Total VAT Amount", "Si", "Depende", "PDF / calculo", "IVA que se usa en Lucy", "", "", ""],
        ["Withholding Tax Amount", "Si/No", "Depende", "PDF / otro", "Retenciones si aplican", "", "", ""],
        ["Total Impuestos Retenidos", "Si/No", "Depende", "PDF / otro", "Si es el mismo campo anterior o no", "", "", ""],
        ["Payment Terms", "Si/No", "Depende", "PDF / maestro vendor", "Codigo o texto exacto", "", "", ""],
        ["Payment Mode", "Si/No", "Depende", "PDF / maestro vendor", "Si Lucy lo usa en este proceso", "", "", ""],
        ["Due Date", "Si/No", "Depende", "PDF / regla", "Fecha de vencimiento real", "", "", ""],
        ["SAP_KOSTL", "Si", "Si", "Control / macro / contratos / otro", "CECO final", "", "", ""],
        ["SAP_PROJK", "Si/No", "Depende", "Macro / otro", "Profit center final", "", "", ""],
        ["Vat Code", "Si", "Depende", "Regla de negocio", "Codigo de IVA en header", "", "", ""],
        ["Calculate Tax", "Si/No", "Depende", "Regla de negocio", "Cuando debe marcarse", "", "", ""],
    ]
    write_table(
        ws,
        1,
        [
            "Campo Header Lucy",
            "Se usa?",
            "Es obligatorio?",
            "Fuente real",
            "Que debe contener",
            "Regla exacta",
            "Si falta o viene ambiguo",
            "Ejemplo real",
        ],
        rows,
        title="Confirmacion de campos de cabecera",
    )
    ws.freeze_panes = "A3"
    autosize(ws, {"A": 28, "B": 12, "C": 14, "D": 24, "E": 26, "F": 30, "G": 28, "H": 22})


def build_allocated_costs(ws):
    rows = [
        ["Posting Key", "Si", "Si", "Macro / SOP / historico", "40, 50 u otro", "Cuando va 40 y cuando 50", "", ""],
        ["GL Account", "Si", "Si", "Chart of accounts / macro / contratos", "Cuenta exacta a postear", "Como se decide RF/RV", "", ""],
        ["Tax Code", "Si", "Depende", "SOP / historico", "VW, VQ, V0, otro", "Cuando aplica cada uno", "", ""],
        ["Cost Center", "Si", "Si", "Control / macro / contratos", "CECO final", "", "", ""],
        ["Profit Center", "Si/No", "Depende", "Macro / otro", "Valor exacto si aplica", "", "", ""],
        ["Text", "Si", "Si", "Regla de negocio", "Texto exacto que se pega en Lucy", "Formato final aprobado", "", ""],
        ["Amount", "Si", "Si", "PDF / calculo", "Valor por linea", "Si usa subtotal, neto, 33%, 50%, etc.", "", ""],
        ["VW VAT line", "Si/No", "Depende", "Prorrateo / historico", "Cuando sale una linea VW", "", "", ""],
        ["VQ VAT line", "Si/No", "Depende", "Prorrateo / historico", "Cuando sale una linea VQ", "", "", ""],
    ]
    write_table(
        ws,
        1,
        [
            "Campo Allocated Costs",
            "Se usa?",
            "Es obligatorio?",
            "Fuente real",
            "Que debe contener",
            "Regla exacta",
            "Si falta o viene ambiguo",
            "Ejemplo real",
        ],
        rows,
        title="Confirmacion de lineas contables",
    )
    ws.freeze_panes = "A3"
    autosize(ws, {"A": 28, "B": 12, "C": 14, "D": 24, "E": 26, "F": 30, "G": 28, "H": 22})


def build_distribution(ws):
    rows = [
        ["33%", "Si/No", "Archivo de distribucion / PDF / otro", "3", "Subtotal, IVA, retenciones, total?", "Si son exactamente iguales o no", "Ejemplo real", ""],
        ["50%", "Si/No", "Archivo de distribucion / PDF / otro", "2", "Subtotal, IVA, retenciones, total?", "Si son exactamente iguales o no", "Ejemplo real", ""],
        ["Factura con varios conceptos", "Si/No", "PDF", "Confirmar", "Si el split depende de conceptos del PDF o solo del archivo de distribucion", "", "", ""],
        ["Vendor destino", "Si", "Archivo de distribucion / historico", "Vendor exacto a usar en cada posteo", "", "", ""],
        ["Text por split", "Si/No", "Regla de negocio", "Si cambia texto entre posteos o no", "", "", ""],
    ]
    write_table(
        ws,
        1,
        [
            "Caso de split",
            "Se usa?",
            "Como se detecta",
            "Cuantos posteos salen",
            "Que se reparte",
            "Regla exacta",
            "Ejemplo real",
            "Observaciones",
        ],
        rows,
        title="Reglas de distribucion y multiples posteos",
    )
    ws.freeze_panes = "A3"
    autosize(ws, {"A": 28, "B": 12, "C": 28, "D": 18, "E": 22, "F": 28, "G": 22, "H": 26})


def build_pdf_formats(ws):
    rows = [
        ["IVA 0%", "Como se reconoce y que campos deben salir", "", "", ""],
        ["Subtotal + IVA claro", "Como se toman subtotal e IVA", "", "", ""],
        ["Con incentivo o descuento", "Que valor se usa en Lucy", "", "", ""],
        ["Con saldo anterior", "Se ignora o se incluye?", "", "", ""],
        ["Con retenciones", "Como se calculan y donde se postean", "", "", ""],
        ["Con varios conceptos/items", "Como se decide split o reparticion", "", "", ""],
        ["Sin fecha clara", "Que se hace si no trae fecha legible", "", "", ""],
        ["Sin IVA claro", "Que se hace si el PDF no separa IVA", "", "", ""],
    ]
    write_table(
        ws,
        1,
        ["Tipo de factura", "Regla real", "Ejemplo PDF", "Salida esperada", "Observaciones"],
        rows,
        title="Tipos de factura y reglas por formato",
    )
    ws.freeze_panes = "A3"
    autosize(ws, {"A": 26, "B": 40, "C": 20, "D": 30, "E": 24})


def build_real_cases(ws):
    rows = [
        ["", "", "", "", "", "", "", "", "", "", "", "", "", "", "", "", "", "", ""],
        ["", "", "", "", "", "", "", "", "", "", "", "", "", "", "", "", "", "", ""],
        ["", "", "", "", "", "", "", "", "", "", "", "", "", "", "", "", "", "", ""],
    ]
    write_table(
        ws,
        1,
        [
            "Invoice_ID",
            "PDF file",
            "Store",
            "Vendor final",
            "CECO final",
            "Rent type",
            "Account final",
            "Subtotal real",
            "IVA real",
            "Retenciones reales",
            "Total real",
            "Split aplica?",
            "Numero de posteos",
            "Header correcto confirmado?",
            "Allocated costs correctos confirmados?",
            "Que se pego en Lucy",
            "Que regla especial aplico",
            "Aprobado por usuario?",
            "Notas",
        ],
        rows,
        title="Casos reales resueltos por el usuario",
    )
    ws.freeze_panes = "A3"
    autosize(ws, {"A": 18, "B": 24, "C": 12, "D": 16, "E": 12, "F": 12, "G": 16, "H": 14, "I": 14, "J": 18, "K": 14, "L": 12, "M": 16, "N": 22, "O": 28, "P": 28, "Q": 28, "R": 18, "S": 24})


def build_blocking_rules(ws):
    rows = [
        ["Falta CECO", "Bloquear factura / generar con flag / bloquear lote", "", ""],
        ["Falta IVA", "Bloquear factura / generar con flag / bloquear lote", "", ""],
        ["Falta Due Date", "Bloquear factura / generar con flag / bloquear lote", "", ""],
        ["Falta Payment Terms", "Bloquear factura / generar con flag / bloquear lote", "", ""],
        ["PDF no legible", "Bloquear factura / generar con flag / bloquear lote", "", ""],
        ["Vendor no encontrado", "Bloquear factura / generar con flag / bloquear lote", "", ""],
        ["Contrato no encontrado", "Bloquear factura / generar con flag / bloquear lote", "", ""],
        ["Split ambiguo", "Bloquear factura / generar con flag / bloquear lote", "", ""],
    ]
    write_table(
        ws,
        1,
        ["Situacion", "Que debe hacer la automatizacion", "Aplica hoy?", "Observaciones"],
        rows,
        title="Reglas de bloqueo o excepcion",
    )
    ws.freeze_panes = "A3"
    autosize(ws, {"A": 26, "B": 42, "C": 14, "D": 26})


def build_open_questions(ws):
    rows = [
        ["El bot debe detenerse si un campo de header viene vacio?", "Pendiente", ""],
        ["El split 33%/50% siempre viene del archivo de distribucion o a veces se deduce por PDF?", "Pendiente", ""],
        ["Invoice Total Amount en header debe ser total factura, total a pagar o subtotal+IVA?", "Pendiente", ""],
        ["Payment Terms debe ser texto del PDF o codigo SAP?", "Pendiente", ""],
        ["Due Date siempre se registra en Lucy o solo en algunos casos?", "Pendiente", ""],
        ["Cuando hay incentivo o saldo anterior, que valor se copia a Lucy exactamente?", "Pendiente", ""],
        ["Que campos del CSV real de headers son obligatorios y cuales se dejan vacios siempre?", "Pendiente", ""],
        ["Cuando usar exactamente la base macro como fuente final y cuando recalcular?", "Pendiente", ""],
    ]
    write_table(
        ws,
        1,
        ["Pregunta", "Estado", "Respuesta final"],
        rows,
        title="Preguntas abiertas para cerrar version productiva",
    )
    ws.freeze_panes = "A3"
    autosize(ws, {"A": 70, "B": 14, "C": 40})


def main():
    wb = Workbook()
    build_instructions(wb.active)
    build_header_fields(wb.create_sheet())
    wb[wb.sheetnames[-1]].title = "02_Header_Fields"
    build_allocated_costs(wb.create_sheet())
    wb[wb.sheetnames[-1]].title = "03_Allocated_Costs"
    build_distribution(wb.create_sheet())
    wb[wb.sheetnames[-1]].title = "04_Distribution"
    build_pdf_formats(wb.create_sheet())
    wb[wb.sheetnames[-1]].title = "05_PDF_Formats"
    build_real_cases(wb.create_sheet())
    wb[wb.sheetnames[-1]].title = "06_Real_Cases"
    build_blocking_rules(wb.create_sheet())
    wb[wb.sheetnames[-1]].title = "07_Blocking_Rules"
    build_open_questions(wb.create_sheet())
    wb[wb.sheetnames[-1]].title = "08_Open_Questions"
    wb.save(OUTPUT_PATH)
    print(f"Created {OUTPUT_PATH}")


if __name__ == "__main__":
    main()
