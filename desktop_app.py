from __future__ import annotations

import ctypes
from pathlib import Path
import sys
from threading import Thread

PROJECT_ROOT = Path(__file__).resolve().parent
SRC_DIR = PROJECT_ROOT / "src"
if str(SRC_DIR) not in sys.path:
    sys.path.insert(0, str(SRC_DIR))

from openpyxl import load_workbook
import webview
from werkzeug.serving import make_server

from lease_accounting.runtime import (
    DICTIONARY_FILENAME,
    LEARNING_DICTIONARY_PATH,
    RESOURCE_DIR,
    cleanup_runtime,
)
from lease_accounting.web import app


APP_TITLE = "EssilorLuxottica - Arriendos Automation"


def _show_error(message: str) -> None:
    ctypes.windll.user32.MessageBoxW(0, message, APP_TITLE, 0x10)


def _validate_runtime() -> list[str]:
    errors = []
    if not LEARNING_DICTIONARY_PATH.exists():
        errors.append(
            f"No se encontro {DICTIONARY_FILENAME} junto al ejecutable.\n\n"
            "Copia el archivo en la misma carpeta que ArriendosAutomation.exe "
            "y vuelve a abrir la aplicacion."
        )
    else:
        try:
            workbook = load_workbook(LEARNING_DICTIONARY_PATH, read_only=True, data_only=True)
            if "Diccionario" not in workbook.sheetnames:
                errors.append("El diccionario no contiene la hoja requerida 'Diccionario'.")
            workbook.close()
        except PermissionError:
            errors.append(
                f"Cierra {DICTIONARY_FILENAME} en Excel antes de abrir la aplicacion."
            )
        except Exception as exc:
            errors.append(f"No se pudo leer {DICTIONARY_FILENAME}: {exc}")

    for relative_path in ("templates/index.html", "static/styles.css"):
        if not (RESOURCE_DIR / relative_path).exists():
            errors.append(f"Falta un recurso interno de la aplicacion: {relative_path}")

    return errors


def _run_smoke_test() -> int:
    errors = _validate_runtime()
    if errors:
        return 2
    with app.test_client() as client:
        response = client.get("/")
        if response.status_code != 200 or b"Arriendos Automation" not in response.data:
            return 3
    return 0


def main() -> int:
    if "--smoke-test" in sys.argv:
        try:
            return _run_smoke_test()
        finally:
            cleanup_runtime()

    errors = _validate_runtime()
    if errors:
        _show_error("\n\n".join(errors))
        cleanup_runtime()
        return 2

    server = make_server("127.0.0.1", 0, app)
    server_thread = Thread(
        target=server.serve_forever,
        name="arriendos-local-server",
        daemon=True,
    )
    server_thread.start()

    webview.settings["ALLOW_DOWNLOADS"] = True
    try:
        webview.create_window(
            APP_TITLE,
            url=f"http://127.0.0.1:{server.server_port}/",
            width=1440,
            height=900,
            min_size=(1024, 700),
            background_color="#080d16",
            text_select=True,
        )
        webview.start(gui="edgechromium", debug=False, private_mode=True)
    finally:
        server.shutdown()
        server.server_close()
        server_thread.join(timeout=3)
        cleanup_runtime()
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
