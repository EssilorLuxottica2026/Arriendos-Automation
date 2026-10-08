import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook, load_workbook

from lease_accounting import web
from lease_accounting.pipeline.processor import PipelineError


class ConceptDictionarySaveTests(unittest.TestCase):
    def setUp(self):
        self.directory = tempfile.TemporaryDirectory()
        self.addCleanup(self.directory.cleanup)
        self.path = Path(self.directory.name) / "dictionary.xlsx"
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Diccionario"
        sheet.append(["Concepto", "GL Account", "Frases/palabras", "Sigla"])
        sheet.append(["ENERGIA", "1532001605", "CONSUMO ENERGIA", "ENERGIA"])
        workbook.save(self.path)
        self.original = self.path.read_bytes()
        self.path_patch = patch.object(web, "LEARNING_DICTIONARY_PATH", self.path)
        self.path_patch.start()
        self.addCleanup(self.path_patch.stop)
        backup_patch = patch.object(web, "_backup_learning_dictionary")
        self.backup = backup_patch.start()
        self.addCleanup(backup_patch.stop)

    def test_encoded_selection_updates_existing_row_without_creating_duplicate(self):
        web._append_learning_dictionary_phrases([
            {"concept": "ENERGIA|1532001605", "phrase": "MANDATO EPM ENERGIA"},
        ])
        workbook = load_workbook(self.path)
        self.addCleanup(workbook.close)
        sheet = workbook["Diccionario"]
        self.assertEqual(sheet.max_row, 2)
        self.assertEqual(sheet["A2"].value, "ENERGIA")
        self.assertEqual(sheet["B2"].value, "1532001605")
        self.assertEqual(sheet["D2"].value, "ENERGIA")
        self.assertIn("MANDATO EPM ENERGIA", sheet["C2"].value)
        self.backup.assert_called_once()

    def test_plain_selection_remains_supported_and_deduplicates_phrases(self):
        web._append_learning_dictionary_phrases([
            {"concept": "ENERGIA", "phrase": "CONSUMO ENERGIA"},
        ])
        workbook = load_workbook(self.path)
        self.addCleanup(workbook.close)
        self.assertEqual(workbook["Diccionario"]["C2"].value, "CONSUMO ENERGIA")

    def test_invalid_selections_do_not_modify_workbook(self):
        for selection in ["ENERGIA|9999999999", "ENERGIA|", "UNKNOWN|1532001605", ""]:
            with self.subTest(selection=selection), self.assertRaises(PipelineError):
                web._append_learning_dictionary_phrases([
                    {"concept": selection, "phrase": "NEW PHRASE"},
                ])
            self.assertEqual(self.path.read_bytes(), self.original)
        self.backup.assert_not_called()

    def test_route_passes_encoded_selection_to_validated_writer(self):
        with (
            patch.object(web, "_load_review_job", return_value={}),
            patch.object(web, "_run_pipeline_from_job", return_value={}),
            patch.object(web, "render_template", return_value="OK"),
            web.app.test_client() as client,
        ):
            response = client.post("/concept-review/test", data={
                "item_count": "1", "concept_0": "ENERGIA|1532001605",
                "phrase_0": "MANDATO EPM ENERGIA",
            })
        self.assertEqual(response.status_code, 200)
        workbook = load_workbook(self.path)
        self.addCleanup(workbook.close)
        self.assertEqual(workbook["Diccionario"].max_row, 2)

    def test_invalid_later_assignment_does_not_save_earlier_assignment(self):
        with self.assertRaises(PipelineError):
            web._append_learning_dictionary_phrases([
                {"concept": "ENERGIA|1532001605", "phrase": "NEW PHRASE"},
                {"concept": "UNKNOWN|1532001605", "phrase": "OTHER PHRASE"},
            ])
        self.assertEqual(self.path.read_bytes(), self.original)
        self.backup.assert_not_called()


if __name__ == "__main__":
    unittest.main()
